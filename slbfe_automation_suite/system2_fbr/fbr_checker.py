#!/usr/bin/env python3
"""
SLBFE - Family Background Report (FBR) Verification Checker   (SYSTEM 2)
=========================================================================
INDEPENDENT SYSTEM - FEMALE APPLICANTS ONLY.

For every FEMALE applicant in name_female.txt it:
  1. Opens the SLBFE 3-method login page and waits for YOUR login, then walks
     the portal menu (main page -> Common Programmes -> Family Background
     Report Verification) to the FBR search page (fbr_sms.asp).
  2. Searches by CURRENT passport number, then by CURRENT NIC number
     (the form accepts one identifier at a time, exactly like the portal).
  3. Validates that the returned row really belongs to that person's
     CURRENT passport / CURRENT NIC (a row for an old number is flagged MISMATCH).
  4. Reads "Date (SMS Received)", adds 6 months and classifies:
     VALID / EXPIRES SOON (<=14 days) / EXPIRED.
  5. Writes one colour-coded Excel workbook with a sheet per outcome group,
     a CSV, fbr_summary.json (for the dashboard), and an evidence screenshot
     whenever an FBR row is found.

FEMALE-ONLY FILTER
  * Rows whose name contains a known Sinhala/Tamil female given-name token are
    included; known male tokens are skipped (the skip list is printed at start).
  * A "Gender" / "Sex" column overrides the name heuristic (only cells starting
    with F are checked). Lines starting with '#' are always ignored.
  * Names the heuristic cannot decide are skipped and listed as UNCERTAIN so
    you can add a Gender column (F) to force-include them.

Every run re-checks every female from scratch (nothing is cached) because the
expiry is recalculated against today's date each time.

This script is fully standalone - it shares NO code, session file or output
with System 1 (the Pre Departure Training checker).

Setup
-----
    1_Install_Dependencies.bat      (or: pip install playwright openpyxl
                                     && python -m playwright install chromium)

Run
---
    python System2_FBR_Checker.py                  # browser opens the 3-method login
    python System2_FBR_Checker.py --only 1-10
    python System2_FBR_Checker.py --only 3,7,15
    python System2_FBR_Checker.py --reuse-session  # skip login if a saved session works
    python System2_FBR_Checker.py --headless       # needs SLBFE_USER/SLBFE_PASS env vars
    python System2_FBR_Checker.py --names name_female.txt --out output

LOGIN  (exact entry URL, never altered: https://services.slbfe.lk/account/ext_account_main)
  By DEFAULT every run opens that page showing all 3 login methods
  (Public Users / Authorized Personnel Only / Development Officer). You click
  'Access' under Authorized Personnel Only and type your credentials; the
  script detects the successful login itself and then walks the portal menu.
  Other options:
    * SLBFE_USER / SLBFE_PASS environment variables -> fully automatic login
      (also works with --headless)
    * --reuse-session -> reuse slbfe_session_fbr.json if it is still valid
  In --headless mode without env vars it prompts for credentials in the terminal.
"""

from __future__ import annotations

import argparse
import calendar
import csv
import getpass
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field, asdict
from datetime import date, datetime
from pathlib import Path

from playwright.sync_api import (
    Page,
    TimeoutError as PWTimeout,
    sync_playwright,
)

# ----------------------------------------------------------------------------
# CONFIG  (all paths stay inside this system's own folder -> full independence)
# ----------------------------------------------------------------------------
BASE = "https://applications.slbfe.lk"
WELCOME_URL = f"{BASE}/logroot/an_welcomenew.asp"
FBR_URL = f"{BASE}/logroot/agency/cmpl/fbr_sms.asp"

# Exact SLBFE login entry URL - opens the page with all 3 login methods
# (Public Users / Authorized Personnel Only / Development Officer).
# Keep this URL exactly as provided - do not deep-link past it.
LOGIN_URL = "https://services.slbfe.lk/account/ext_account_main"
# "Authorized Personnel Only" card -> its "Access" link (relative href on the page)
AUTH_ACCESS = 'a[href$="ext_account_"]'
LOGIN_USER = 'input[name="usrLoginName"]'
LOGIN_PASS = 'input[name="usrPassWord"]'
LOGIN_SUBMIT = 'button[type="submit"]'

VALIDITY_MONTHS = 6            # FBR is valid for 6 months from SMS-received date
EXPIRES_SOON_DAYS = 14
NAV_TIMEOUT = 60_000           # ms
DELAY_BETWEEN = 1.0            # seconds between searches (gentle pacing)
RELOGIN_RETRIES = 2            # extra full re-logins after a mid-run session expiry

# System 2 keeps its OWN session file - never shares System 1's login state.
SCRIPT_DIR = Path(__file__).resolve().parent
SESSION_FILE = SCRIPT_DIR / "slbfe_session_fbr.json"
DEFAULT_NAMES = SCRIPT_DIR / "name_female.txt"
DEFAULT_OUTDIR = SCRIPT_DIR / "output"

USER_ENV, PASS_ENV = "SLBFE_USER", "SLBFE_PASS"

# ----------------------------------------------------------------------------
# FEMALE-ONLY NAME FILTER
# Tokens are matched as whole words against the space-separated uppercase name.
# A name is female only when it carries a female token and no male token.
# ----------------------------------------------------------------------------
MALE_TOKENS = set("""LAKMAL THARMASEELAN LAHIRU MADURANGA SACHINTHA PRAMAL RANDUNU JAYASHAN USHAN THIWANKA
HESHAN MADHUWANTHA SANKALPA MALSHAN MADUSHAN KAVINDA SANDARUWAN SANDUN CHAMARA SUMUDITHA SURESH ISHAN MADHUSHA
CHANAKA ISHARA BANDARUWAN RASHMIKA ASHITHA RANGA THARINDU CHAMIKARA RAVIN PRAMANGA SURAJ PASINDU KUMARA
DULAKSHANA JAYASHANTHA HASUN LAKSHAN NAVEEN PRIYADARSHANA CHANUKA DULSHAN GAYAN PRADEEP DUMIDU PIYANKARA
DAMITH PRASAD KOSHALA DANANJAYA GIHAN RUKMAL CHRISTIAN DIMUTHU NUWAN SRILAL THEEKSHANA NALAKA MIHINTHAL
EMMANUWEL SENURA DIMAL HARITH NAWIN SUDESH PRIYANTHA NISHANTHA SOVIS ANTANI LEHAN KUMAR PRAKASH VIRAJ
RUCHIRA KAUSHAN RASHINTHA GIMES KALHARA KAVEESHA GEESHAN SHANAKA CRISHAN DHANANJAYA DILRUK ROMAL NIDUKTHA
KASUN JAYASANKA CHAMATH CHAMIL SHIWANKA ASHAN THARUKA MILINDA RAJIVE SANJAYA PUBUDITHA SENAHAS RANIDU
SHEHARA BOSHAN RUPAS ANTON DUSHMINDA DILSHAN CHAMOD UDAYANGA BINAL IROSHAN SANJEEWA NIRASH SUPUN SANJALA
NIDUSHAN LAKSHITHA LAKSHEN CHAMITH NAGESWARAN BALASUBRAMANIYAM CHANDEEPA NISADA""".split())

FEMALE_TOKENS = set("""THARUSHIKA NADEESHANI DILANI DELKA CHATHURANGANI NISHADINI PRASANGIKA SHASHINI NIMASHA SEMINA
GANGANATHA SEWWANDI THAKSHILA NETHMINI MADHUSHANI THARAKA NIRASHA HIRUNI SALANI SAMADHI SHASHIPRABHA SANSALA
THARUSHI DAMMIKA PRIYADARSHANI SANDUNI MADHUSHI PIYUMI IMASHI MALKI NETHULI PRAVINDAYA KAVISHI SANJANA
CHANCHALA KAVINDI NETHMIKA MANUSHA IMALI HANSIKA SHEHANI NISANSA SAVINDI NISANSALA SHERONI DINITHI KAVINDYA
JAYALATH SENARA CHATURIKA DILMI SHALANI CRISENTA NIMESHA PRIYASANI DIMITHRIKA SHIROMI SUDESHIKA LAKSHANI
PUNYALATHA RUVINDI DEVI NIPUNEE KAWSHALYA RUWANI NIDISHANI SUNALI SASHIKA GAYATHRI KUMUDUNI NIPUNI SADAMINI
SEWMINI SANDALI HIRUSHI MADHAVI NAWODYA OSHADI RASHENI NAYANJANA RAPINA PIUMA WASANA KUMARI KAWYA
MADHUWANTHI NITHISHA RISHMI VISHMI MADARA SHAVINDYA GANANJANI NATHASHA NATASHA SEHANI HIMASHA SIWMINI NEELIYA
NADEERA CHATHURI MANISHA DILUKSHI IRESHA KANCHANAMALA DINESHA BIMALI YASHODHA MALSHA RASHINI PIYUMALI GIMHANI
RENUKA DAMAYANTHI SRI""".split())


class SessionExpiredError(RuntimeError):
    """The ASP session died and a fresh manual login is required."""


# ----------------------------------------------------------------------------
# DATA MODELS
# ----------------------------------------------------------------------------
@dataclass
class Applicant:
    no: int
    name: str
    passport: str
    nic: str
    prev_passport: str = ""
    prev_nic: str = ""


@dataclass
class FbrRow:
    sms_dt: datetime
    passport: str
    nic: str
    agency_lno: str
    status: str


@dataclass
class Result:
    no: int
    name: str
    current_passport: str
    current_nic: str
    searched_by: str = ""
    match_on: str = ""             # which CURRENT identifier matched
    fbr_date: str = ""
    expiry_date: str = ""
    days_left: str = ""
    expiry_state: str = ""         # VALID / EXPIRED / EXPIRES SOON
    fbr_status: str = ""           # Recommended / Pending ...
    result_passport: str = ""
    result_nic: str = ""
    agency_lno: str = ""
    validation: str = ""           # OK / CHECK / MISMATCH / NOT FOUND / ERROR
    remark: str = ""
    screenshot: str = ""
    checked_at: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S"))


# ----------------------------------------------------------------------------
# INPUT: female-only parsing of the markdown applicant table
# ----------------------------------------------------------------------------
def _clean(v: str) -> str:
    v = v.strip().strip("*").strip().replace("\\-", "-")
    return "" if v in ("-", "") else v.strip()


def classify_gender(name: str) -> str:
    """Return 'F', 'M' or '?' using whole-word given-name tokens."""
    words = set(name.upper().split())
    male = bool(words & MALE_TOKENS)
    female = bool(words & FEMALE_TOKENS)
    if female and not male:
        return "F"
    if male and not female:
        return "M"
    return "?"


def load_applicants_with_skips(path: str) -> tuple[list[Applicant], list[tuple[int, str, str]], list[tuple[int, str]]]:
    """Return (female_applicants, skipped[(no, name, reason)], uncertain[(no, name)]).

    A Gender/Sex column (value starting with F) overrides the name heuristic.
    """
    out: list[Applicant] = []
    skipped: list[tuple[int, str, str]] = []
    uncertain: list[tuple[int, str]] = []
    gender_col = None
    for line in Path(path).read_text(encoding="utf-8", errors="ignore").splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) < 4:
            continue
        # remember an optional Gender/Sex column if the table ever has one
        if "gender" in cells[0].lower() or any("gender" in c.lower() or c.lower() == "sex" for c in cells):
            for i, c in enumerate(cells):
                if "gender" in c.lower() or c.lower() == "sex":
                    gender_col = i
            continue
        if not cells[0].isdigit():
            continue  # separator / other rows
        cells += [""] * max(0, 7 - len(cells))
        no = int(cells[0])
        name = _clean(cells[1])

        explicit = cells[gender_col].strip().upper() if gender_col is not None and gender_col < len(cells) else ""
        if explicit:
            is_female = explicit.startswith("F")
            why = f"Gender column = {cells[gender_col].strip()}"
            g = "F" if is_female else "M"
        else:
            g = classify_gender(name)
            is_female = g == "F"
            why = "male name token" if g == "M" else "name not recognised as female"

        if g == "?":
            uncertain.append((no, name))   # reported separately, never double-counted
            continue
        if not is_female:
            skipped.append((no, name, why))
            continue

        out.append(
            Applicant(
                no=no,
                name=name,
                passport=_clean(cells[2]).upper().replace(" ", ""),
                nic=_clean(cells[3]).upper().replace(" ", ""),
                prev_passport=_clean(cells[4]).upper().replace(" ", ""),
                prev_nic=_clean(cells[5]).upper().replace(" ", ""),
            )
        )
    return out, skipped, uncertain


def load_applicants(path: str) -> list[Applicant]:
    """Backward-compatible wrapper used by the unit tests."""
    return load_applicants_with_skips(path)[0]


def select_applicants(all_: list[Applicant], spec: str | None) -> list[Applicant]:
    if not spec:
        return all_
    wanted: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            wanted.update(range(int(a), int(b) + 1))
        elif part:
            wanted.add(int(part))
    return [a for a in all_ if a.no in wanted]


# ----------------------------------------------------------------------------
# DATE LOGIC
# ----------------------------------------------------------------------------
def add_months(d: date, months: int) -> date:
    """Calendar-month addition with end-of-month clamping (31 Aug + 6m -> Feb 28/29)."""
    idx = d.month - 1 + months
    y, m = d.year + idx // 12, idx % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def expiry_info(sms_dt: datetime, today: date | None = None) -> tuple[date, int, str]:
    today = today or date.today()
    expiry = add_months(sms_dt.date(), VALIDITY_MONTHS)
    days_left = (expiry - today).days
    if days_left < 0:
        state = "EXPIRED"
    elif days_left <= EXPIRES_SOON_DAYS:
        state = "EXPIRES SOON"
    else:
        state = "VALID"
    return expiry, days_left, state


# ----------------------------------------------------------------------------
# LOGIN / SESSION
# ----------------------------------------------------------------------------
def fbr_form_present(page: Page) -> bool:
    try:
        return page.locator('input[name="ppno"]').count() > 0
    except Exception:
        return False


def session_expired(page: Page) -> bool:
    url = (page.url or "").lower()
    if "wrongway" in url or "gohome" in url:
        return True
    try:
        return "session has expired" in page.inner_text("body").lower()
    except Exception:
        return False


def _login_screen_visible(pg: Page) -> bool:
    """True while the browser is still on an SLBFE login page:
    the 3-method page (has the Access link) or the agency Sign In form."""
    try:
        if "services.slbfe.lk/account" not in (pg.url or "").lower():
            return False
        return pg.locator(AUTH_ACCESS).count() > 0 or pg.locator(LOGIN_USER).count() > 0
    except Exception:
        return False


def _any_page_has_wrongway(context) -> bool:
    for pg in context.pages:
        try:
            if "wrongway" in (pg.url or "").lower():
                return True
        except Exception:
            pass
    return False


def _verify_session(page: Page, context) -> bool:
    """Probe the FBR page. Works no matter which tab/window completed the
    login, because the ASP session cookie is shared browser-context-wide."""
    try:
        page.goto(FBR_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
        if fbr_form_present(page):
            context.storage_state(path=str(SESSION_FILE))
            return True
    except Exception:
        pass
    return False


def _submit_login_form(page: Page, user: str, password: str) -> tuple[bool, str]:
    """Open the 3-method login page, click 'Authorized Personnel Only -> Access',
    fill the Sign In form (posts to applications.slbfe.lk/chkuser.asp), and
    verify the resulting session against the FBR page."""
    context = page.context
    page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
    page.wait_for_selector(AUTH_ACCESS, timeout=NAV_TIMEOUT)
    page.click(AUTH_ACCESS)
    page.wait_for_selector(LOGIN_USER, timeout=NAV_TIMEOUT)
    page.fill(LOGIN_USER, user)
    page.fill(LOGIN_PASS, password)
    try:
        page.check("#checkRememberPassword")   # keep the session alive longer
    except Exception:
        pass
    page.click(LOGIN_SUBMIT)

    # wait until we leave the Sign In form or a wrongway rejection appears
    deadline = time.time() + 30
    while time.time() < deadline:
        if ("wrongway" in (page.url or "").lower()
                or "services.slbfe.lk/account" not in (page.url or "").lower()
                or _any_page_has_wrongway(context)):
            break
        page.wait_for_timeout(500)

    if "wrongway" in (page.url or "").lower() or _any_page_has_wrongway(context):
        try:
            msg = re.sub(r"\s+", " ", page.inner_text("body")).strip()
        except Exception:
            msg = "Wrong User Name or Password"
        return False, msg[:140]

    if _verify_session(page, context):
        return True, ""
    return False, "Login did not produce a valid portal session."


def _prompt_credentials() -> tuple[str, str]:
    user = input("SLBFE username: ").strip()
    password = getpass.getpass("SLBFE password: ")
    return user, password


def _wait_manual_browser_login(page: Page) -> None:
    """Headed login: the browser shows the 3-method SLBFE page and the user
    signs in manually. The script polls the shared browser session (works even
    if the portal opens in a new tab) and continues automatically."""
    context = page.context
    popups = []
    context.on("page", lambda p: popups.append(p))

    page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
    print("\n" + "=" * 70)
    print("LOGIN REQUIRED - the browser now shows the SLBFE page with 3 login methods.")
    print("1. Click 'Access' under 'Authorized Personnel Only' (for agencies).")
    print("2. Enter your username and password on the Sign In page.")
    print("(Credentials are typed into the website, not this terminal.)")
    print("The script continues automatically after a successful login.")
    print("=" * 70)

    deadline = time.time() + 300  # 5 minutes
    last_note = 0.0
    while time.time() < deadline:
        active = popups[-1] if popups and not popups[-1].is_closed() else page
        try:
            active.wait_for_load_state("domcontentloaded", timeout=3000)
        except Exception:
            pass

        if "wrongway" in (active.url or "").lower() or _any_page_has_wrongway(context):
            print("[login] sign-in was rejected - returning to the 3-method login page...")
            try:
                page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
            except Exception:
                pass
            time.sleep(1)
            continue

        if _login_screen_visible(active):
            if time.time() - last_note > 15:
                print("[login] waiting for you to complete the sign-in in the browser...")
                last_note = time.time()
            time.sleep(2)
            continue

        # left the login pages -> check whether the session cookie works
        print("[login] sign-in detected - verifying the portal session...")
        if _verify_session(page, context):
            for p in popups:   # close any helper tabs, keep the working page
                try:
                    if not p.is_closed():
                        p.close()
                except Exception:
                    pass
            print("[login] login OK - session saved.")
            return
        time.sleep(2)

    raise SystemExit("Login was not completed within 5 minutes. Re-run the script.")


def enter_fbr_section(page: Page) -> bool:
    """Walk the portal the way a user does: main welcome page -> the
    'Family Background Report Verification' menu link -> the FBR search form.
    Falls back to opening fbr_sms.asp directly if the menu cannot be found."""
    link_re = re.compile(r"family\s+background", re.I)
    via_menu = False
    try:
        page.goto(WELCOME_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
        for frame in page.frames:
            try:
                loc = frame.get_by_text(link_re)
                if loc.count() > 0:
                    try:
                        with page.expect_navigation(wait_until="domcontentloaded", timeout=NAV_TIMEOUT):
                            loc.first.click()
                    except PWTimeout:
                        page.wait_for_load_state("domcontentloaded", timeout=10000)
                    via_menu = True
                    break
            except Exception:
                continue
    except Exception:
        pass

    if not fbr_form_present(page):
        try:
            page.goto(FBR_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
        except Exception:
            pass
    ok = fbr_form_present(page)
    if ok:
        print(f"[nav] Family Background Report Verification page opened "
              f"({'via portal menu' if via_menu else 'direct link fallback'}) - automation starting.")
    else:
        print("[nav] WARNING: could not open the FBR search page.")
    return ok


def ensure_login(page: Page, headless: bool, allow_saved: bool = False) -> None:
    """Authenticate, then park the browser on the FBR search page.

    allow_saved=False (default): in a VISIBLE browser ALWAYS show the 3-method
    login page first. allow_saved=True (--reuse-session) also accepts a still-
    valid saved session. In --headless mode there is no window to type into,
    so an existing session is probed first either way.
    Environment credentials always take over when SLBFE_USER/SLBFE_PASS exist.
    """
    if allow_saved or headless:
        page.goto(FBR_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
        if fbr_form_present(page):
            page.context.storage_state(path=str(SESSION_FILE))
            print("[login] saved session is valid")
            enter_fbr_section(page)
            return

    # environment variables -> fully automatic login through the 3-method page
    user, password = os.getenv(USER_ENV), os.getenv(PASS_ENV)
    if user and password:
        ok, msg = _submit_login_form(page, user, password)
        if ok:
            print("[login] automatic login OK (environment credentials)")
            enter_fbr_section(page)
            return
        print(f"[login] environment credentials rejected: {msg or 'unknown error'}")

    # visible browser: the 3-method login page stays until the user signs in
    if not headless:
        _wait_manual_browser_login(page)
        enter_fbr_section(page)
        return

    # headless: no browser to type into -> terminal prompt (3 attempts)
    print("[login] no valid session and no SLBFE_USER/SLBFE_PASS environment variables.")
    for attempt in range(3):
        try:
            user, password = _prompt_credentials()
        except (EOFError, KeyboardInterrupt):
            break
        if not user:
            break
        ok, msg = _submit_login_form(page, user, password)
        if ok:
            print("[login] login OK - session saved")
            enter_fbr_section(page)
            return
        print(f"[login] rejected ({msg or 'unknown error'}); {2 - attempt} attempt(s) left.")
    raise SessionExpiredError(
        "Headless login unavailable. Set SLBFE_USER/SLBFE_PASS or run without --headless."
    )


# ----------------------------------------------------------------------------
# SEARCH + PARSE
# ----------------------------------------------------------------------------
DATE_RE = re.compile(r"(\d{2})/(\d{2})/(\d{4})(?:\s*-?\s*(\d{2}:\d{2}:\d{2}))?")


def parse_result_rows(page: Page) -> list[FbrRow]:
    """Reads the result table (columns: #, Date, Passport, NIC, Agency LNO, Status)."""
    raw = page.evaluate(
        """() => {
            const rows = [];
            document.querySelectorAll('tr').forEach(tr => {
                const tds = Array.from(tr.children).filter(c => c.tagName === 'TD');
                // innermost data rows only: >=5 cells, no nested table
                if (tds.length >= 5 && !tr.querySelector('table')) {
                    rows.push(tds.map(td => td.innerText.replace(/\\s+/g,' ').trim()));
                }
            });
            return rows;
        }"""
    )
    parsed: list[FbrRow] = []
    for cells in raw:
        joined = " | ".join(cells)
        if not DATE_RE.search(joined):
            continue
        di = next((i for i, c in enumerate(cells) if DATE_RE.search(c)), None)
        if di is None or len(cells) < di + 5:
            continue
        dm = DATE_RE.search(cells[di])
        d, mo, y, t = dm.groups()
        hh, mm, ss = (t or "00:00:00").split(":")
        parsed.append(
            FbrRow(
                sms_dt=datetime(int(y), int(mo), int(d), int(hh), int(mm), int(ss)),
                passport=cells[di + 1].upper().replace(" ", ""),
                nic=cells[di + 2].upper().replace(" ", ""),
                agency_lno=cells[di + 3],
                status=cells[di + 4].strip(),
            )
        )
    return parsed


def search(page: Page, passport: str = "", nic: str = "") -> list[FbrRow]:
    """Opens the form fresh, fills ONE identifier, submits, parses the table."""
    page.goto(FBR_URL, wait_until="domcontentloaded", timeout=NAV_TIMEOUT)
    if session_expired(page) or not fbr_form_present(page):
        raise SessionExpiredError("FBR search form missing - session expired.")

    page.fill('input[name="ppno"]', passport)
    page.fill('input[name="idno"]', nic)
    try:
        with page.expect_navigation(timeout=NAV_TIMEOUT):
            page.click('input[name="search"]')
    except PWTimeout:
        # the site's chk() alert can block submit: we never left the form page
        if fbr_form_present(page):
            return []
        page.wait_for_load_state("domcontentloaded", timeout=NAV_TIMEOUT)

    if session_expired(page):
        raise SessionExpiredError("Session expired after submitting the search.")
    return parse_result_rows(page)


def search_with_retry(page: Page, **kw) -> list[FbrRow]:
    last: Exception | None = None
    for attempt in range(2):
        try:
            return search(page, **kw)
        except SessionExpiredError:
            raise
        except Exception as e:  # noqa: BLE001
            last = e
            time.sleep(2 * (attempt + 1))
    raise last  # type: ignore[misc]


def safe_filename(text: str, limit: int = 80) -> str:
    text = re.sub(r"[^A-Za-z0-9 _\-]", "", text).strip()
    return re.sub(r"\s+", "_", text)[:limit]


# ----------------------------------------------------------------------------
# PER-APPLICANT VALIDATION
# ----------------------------------------------------------------------------
def check_applicant(page: Page, a: Applicant, today: date, shot_dir: Path, delay: float) -> Result:
    res = Result(no=a.no, name=a.name, current_passport=a.passport, current_nic=a.nic)

    attempts: list[tuple[str, dict]] = []
    if a.passport:
        attempts.append(("Passport", {"passport": a.passport}))
    if a.nic:
        attempts.append(("NIC", {"nic": a.nic}))

    seen: list[tuple[str, FbrRow]] = []
    notes: list[str] = []
    try:
        for label, kw in attempts:
            try:
                rows = search_with_retry(page, **kw)
            except SessionExpiredError:
                raise
            except Exception as e:  # noqa: BLE001 - one bad identifier must not kill the NIC fallback
                notes.append(f"{label} search error: {type(e).__name__}")
                continue
            for r in rows:
                seen.append((label, r))
            if not rows:
                notes.append(f"No row for {label}")
            time.sleep(delay)
    except SessionExpiredError:
        raise
    except Exception as e:  # noqa: BLE001
        res.validation, res.remark = "ERROR", f"{type(e).__name__}: {e}"[:200]
        return res

    # Keep only rows belonging to the CURRENT passport or CURRENT NIC
    valid: list[tuple[str, FbrRow, str]] = []
    stray: list[FbrRow] = []
    for label, r in seen:
        pp_ok = bool(a.passport) and r.passport == a.passport
        id_ok = bool(a.nic) and r.nic == a.nic
        if pp_ok or id_ok:
            matched = "Passport+NIC" if (pp_ok and id_ok) else ("Passport" if pp_ok else "NIC")
            valid.append((label, r, matched))
        else:
            stray.append(r)

    if not valid:
        if stray:
            s = stray[0]
            res.validation = "MISMATCH"
            res.result_passport, res.result_nic = s.passport, s.nic
            res.remark = (
                "Row returned but passport/NIC differ from CURRENT ones "
                f"(returned {s.passport} / {s.nic}); possibly an old passport or ID."
            )
        else:
            res.validation = "NOT FOUND"
            res.remark = "; ".join(notes) or "No FBR record"
        return res

    # Newest report wins across the two searches
    label, row, matched = max(valid, key=lambda t: t[1].sms_dt)
    expiry, days_left, state = expiry_info(row.sms_dt, today)

    old_id = []
    if a.passport and row.passport != a.passport:
        old_id.append(f"report filed under passport {row.passport} (current is {a.passport})")
    if a.nic and row.nic != a.nic:
        old_id.append(f"report filed under NIC {row.nic} (current is {a.nic})")

    res.searched_by = label
    res.match_on = matched
    res.fbr_date = row.sms_dt.strftime("%d/%m/%Y %H:%M:%S")
    res.expiry_date = expiry.strftime("%d/%m/%Y")
    res.days_left = str(days_left)
    res.expiry_state = state
    res.fbr_status = row.status
    res.result_passport, res.result_nic = row.passport, row.nic
    res.agency_lno = row.agency_lno
    res.validation = "CHECK" if old_id else "OK"
    parts = list(old_id)
    if state == "EXPIRED":
        parts.append(f"Report expired {-days_left} day(s) ago")
    elif state == "EXPIRES SOON":
        parts.append(f"Expires in {days_left} day(s)")
    res.remark = " | ".join(parts)

    # evidence screenshot of the result page
    try:
        shot = shot_dir / f"{a.no:03d}_{safe_filename(a.name)}.png"
        page.screenshot(path=str(shot), full_page=True)
        res.screenshot = str(shot.relative_to(SCRIPT_DIR))
    except Exception:
        pass
    return res


# ----------------------------------------------------------------------------
# OUTPUT
# ----------------------------------------------------------------------------
COLUMNS = [
    ("No", "no"), ("Name", "name"),
    ("Current Passport", "current_passport"), ("Current NIC", "current_nic"),
    ("Validation", "validation"), ("Expiry State", "expiry_state"),
    ("FBR Date (SMS Received)", "fbr_date"), ("Expiry Date (+6 months)", "expiry_date"),
    ("Days Left", "days_left"), ("FBR Status", "fbr_status"),
    ("Matched On", "match_on"), ("Returned Passport", "result_passport"),
    ("Returned NIC", "result_nic"), ("Agency LNO", "agency_lno"),
    ("Screenshot", "screenshot"), ("Remark", "remark"), ("Checked At", "checked_at"),
]

# workbook sheet -> predicate(result). Every result lands in exactly one group:
# rows with a report (OK/CHECK) follow their expiry state; the rest fall through.
SHEET_GROUPS = [
    ("VALID FBR", lambda r: r.expiry_state == "VALID"),
    ("EXPIRED & EXPIRING", lambda r: r.expiry_state in ("EXPIRED", "EXPIRES SOON")),
    ("NO FBR", lambda r: r.validation == "NOT FOUND"),
    ("REVIEW & ERRORS", lambda r: r.validation in ("MISMATCH", "ERROR") or r.validation == ""),
]


def save_outputs(results: list[Result], outdir: Path) -> tuple[Path, Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_path = outdir / f"fbr_report_{stamp}.csv"
    xlsx_path = outdir / f"fbr_report_{stamp}.xlsx"

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in COLUMNS])
        for r in results:
            d = asdict(r)
            w.writerow([d[k] for _, k in COLUMNS])
    # fixed-name copies so the dashboard / shortcuts always find the latest run
    latest_csv = outdir / "fbr_report_latest.csv"
    latest_csv.write_text(csv_path.read_text(encoding="utf-8-sig"), encoding="utf-8-sig")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        red = PatternFill("solid", fgColor="F8CBAD")
        amber = PatternFill("solid", fgColor="FFE699")
        green = PatternFill("solid", fgColor="C6E0B4")
        grey = PatternFill("solid", fgColor="D9D9D9")
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="1F3864")

        def row_fill(r: Result) -> PatternFill:
            if r.expiry_state == "EXPIRED":
                return red
            if r.expiry_state == "EXPIRES SOON":
                return amber
            if r.expiry_state == "VALID" and r.validation == "OK":
                return green
            return grey

        def populate(ws, rows: list[Result], colour: bool) -> None:
            ws.append([c[0] for c in COLUMNS])
            for cell in ws[1]:
                cell.font = head_font
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center")
            for r in rows:
                d = asdict(r)
                ws.append([d[k] for _, k in COLUMNS])
                if colour:
                    fill = row_fill(r)
                    for cell in ws[ws.max_row]:
                        cell.fill = fill
                    if r.expiry_state == "EXPIRED":
                        for cell in ws[ws.max_row]:
                            cell.font = Font(bold=True, color="9C0006")
            for i, (title, _) in enumerate(COLUMNS, 1):
                width = max([len(title)] + [len(str(c.value or "")) for c in ws[get_column_letter(i)]])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions

        wb = Workbook()
        ws_all = wb.active
        ws_all.title = "All Females"
        populate(ws_all, results, colour=True)
        for sheet_name, pred in SHEET_GROUPS:
            ws = wb.create_sheet(sheet_name[:31])
            populate(ws, [r for r in results if pred(r)], colour=True)
        wb.save(xlsx_path)
        wb.save(outdir / "fbr_report_latest.xlsx")
        print(f"Excel : {xlsx_path}")
        print("        sheets: " + ", ".join(wb.sheetnames))
    except ImportError:
        xlsx_path = Path()
        print("openpyxl not installed -> Excel skipped (pip install openpyxl)")
    print(f"CSV   : {csv_path}")
    return csv_path, xlsx_path


def save_summary(results: list[Result], outdir: Path, total_loaded: int,
                 skipped: list[tuple[int, str, str]], uncertain: list[tuple[int, str]]) -> Path:
    """Compact JSON consumed by the common dashboard."""
    counts: dict[str, int] = {}
    for r in results:
        key = r.expiry_state or r.validation
        counts[key] = counts.get(key, 0) + 1
    summary = {
        "system": "System 2 - Female FBR Verification",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "applicants_in_file": total_loaded,
        "females_checked": len(results),
        "skipped_non_female": [{"no": n, "name": nm, "reason": why} for n, nm, why in skipped],
        "uncertain_names": [{"no": n, "name": nm} for n, nm in uncertain],
        "counts": counts,
        "reports": {
            "csv": "output/fbr_report_latest.csv",
            "xlsx": "output/fbr_report_latest.xlsx",
        },
        "results": [asdict(r) for r in results],
    }
    path = outdir / "fbr_summary.json"
    path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def print_summary(results: list[Result]) -> None:
    cnt = lambda f: sum(1 for r in results if f(r))  # noqa: E731
    valid = [r for r in results if r.expiry_state == "VALID"]
    expiring = [r for r in results if r.expiry_state == "EXPIRES SOON"]
    expired = [r for r in results if r.expiry_state == "EXPIRED"]
    no_fbr = [r for r in results if r.validation == "NOT FOUND"]
    review = [r for r in results if r.validation in ("MISMATCH", "CHECK", "ERROR")]

    print("\n" + "=" * 60)
    print(f"FEMALES CHECKED   : {len(results)}")
    print(f"VALID FBR         : {len(valid)}")
    print(f"EXPIRES SOON      : {len(expiring)}")
    print(f"EXPIRED           : {len(expired)}")
    print(f"NOT FOUND (no FBR): {len(no_fbr)}")
    print(f"REVIEW / ERRORS   : {len(review)}")

    def dump(title: str, rows: list[Result], extra) -> None:
        if not rows:
            return
        print(f"\n--- {title} ({len(rows)}) ---")
        for r in rows:
            print(f"  #{r.no:<3} {r.name[:48]:<48} {extra(r)}")

    dump("GOT VALID FBR", valid, lambda r: f"FBR {r.fbr_date[:10]}  until {r.expiry_date}  [{r.fbr_status}]")
    dump("EXPIRES SOON", expiring, lambda r: f"expires {r.expiry_date} ({r.days_left} d)")
    dump("FBR EXPIRED", expired, lambda r: f"FBR {r.fbr_date[:10]}  expired {r.expiry_date}")
    dump("NO FBR RECORD", no_fbr, lambda r: r.remark[:40])
    dump("REVIEW / ERRORS", review, lambda r: f"{r.validation}: {r.remark[:40]}")
    print("=" * 60)


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description="SLBFE female-only FBR verification + 6-month expiry checker")
    ap.add_argument("--names", default=str(DEFAULT_NAMES), help="female applicant table")
    ap.add_argument("--outdir", default=str(DEFAULT_OUTDIR), help="output folder")
    ap.add_argument("--only", help="e.g. 1-10 or 3,7,15 (applicant numbers)")
    ap.add_argument("--reuse-session", action="store_true",
                    help="reuse slbfe_session_fbr.json instead of showing the login page every run")
    ap.add_argument("--delay", type=float, default=DELAY_BETWEEN, help="seconds between searches")
    ap.add_argument("--headless", action="store_true")
    args = ap.parse_args()

    if not Path(args.names).exists():
        sys.exit(f"Input file not found: {args.names}")
    applicants, skipped, uncertain = load_applicants_with_skips(args.names)
    total_rows = len(applicants) + len(skipped) + len(uncertain)
    if skipped:
        print(f"Female-only filter: skipped {len(skipped)} non-female row(s):")
        for no, nm, why in skipped:
            print(f"   #{no:<3} {nm[:50]:<50} ({why})")
    if uncertain:
        print(f"WARNING: {len(uncertain)} name(s) cannot be gender-identified and were SKIPPED.")
        print("         Add a 'Gender' column with F to include, or correct the name:")
        for no, nm in uncertain:
            print(f"   #{no:<3} {nm}")
    applicants = select_applicants(applicants, args.only)
    if not applicants:
        sys.exit("No female applicants to check - fix the input file / --only value.")
    outdir = Path(args.outdir)
    shot_dir = outdir / "fbr_screenshots"
    shot_dir.mkdir(parents=True, exist_ok=True)
    print(f"\nLoaded {len(applicants)} FEMALE applicant(s) to check "
          f"({len(skipped)} non-female skipped, {len(uncertain)} uncertain). "
          f"Today = {date.today():%d/%m/%Y}")

    results: list[Result] = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=args.headless)
        ctx_kw = {"storage_state": str(SESSION_FILE)} if SESSION_FILE.exists() else {}
        ctx = browser.new_context(viewport={"width": 1400, "height": 900}, **ctx_kw)
        ctx.set_default_timeout(NAV_TIMEOUT)
        page = ctx.new_page()

        # the portal pops alert() on short passport numbers - never let it hang
        page.on("dialog", lambda d: d.accept())

        # default: the 3-method login page appears every run
        ensure_login(page, headless=args.headless, allow_saved=args.reuse_session)
        today = date.today()
        try:
            for i, a in enumerate(applicants, 1):
                r: Result | None = None
                for attempt in range(RELOGIN_RETRIES + 1):
                    try:
                        r = check_applicant(page, a, today, shot_dir, args.delay)
                        break
                    except SessionExpiredError as e:
                        print("  [session] expired - full re-login required...")
                        try:
                            ctx.clear_cookies()
                            ensure_login(page, headless=args.headless, allow_saved=False)
                        except SessionExpiredError:
                            pass  # headless/no creds: fall through to ERROR below
                if r is None:
                    r = Result(no=a.no, name=a.name, current_passport=a.passport,
                               current_nic=a.nic, validation="ERROR",
                               remark="Session expired repeatedly despite re-login - re-run to retry")
                results.append(r)
                flag = r.expiry_state or r.validation
                print(f"[{i:>3}/{len(applicants)}] #{a.no:<3} {a.passport:<10} {a.nic:<13} -> "
                      f"{flag:<13} {r.fbr_status:<12} {r.fbr_date[:10] if r.fbr_date else '-':<10} {r.remark[:50]}")
        except KeyboardInterrupt:
            print("\nInterrupted - saving what has been checked so far...")
        finally:
            browser.close()

    save_outputs(results, outdir)
    save_summary(results, outdir, total_rows, skipped, uncertain)
    print_summary(results)


if __name__ == "__main__":
    main()
