"""
End-to-end tests for fbr_checker.py (System 2) against the local mock FBR
server. No real SLBFE credentials are used.

Covers:
  * female-only filtering (name tokens, explicit Gender column, uncertain skip)
  * portal-menu walk from the welcome page to the FBR search form
  * VALID / EXPIRES SOON / EXPIRED classification, NOT FOUND, MISMATCH,
    short-passport alert then NIC fallback
  * session-expiry -> full re-login retries -> ERROR row, run never crashes
  * multi-sheet Excel workbook, CSV, JSON summary and evidence screenshots
"""
import builtins
import csv
import io
import json
import shutil
import sys
import threading
from contextlib import redirect_stdout
from datetime import date, timedelta
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import fbr_checker as fbr
from mock_fbr_server import Handler

MOCK_BASE = "http://127.0.0.1:8766"
MOCK_FORM = f"{MOCK_BASE}/logroot/agency/cmpl/fbr_sms.asp"
MOCK_WELCOME = f"{MOCK_BASE}/logroot/an_welcomenew.asp"
WORK = ROOT / "tests" / "fbr_workspace"

# Female rows carry real female given-name tokens; the male and unknown rows
# must be removed by the female-only filter before any browser work happens.
NAME_TXT = """# SYSTEM 2 (FBR) - FEMALE APPLICANTS ONLY
| No. | Name | Current Passport Number | Current ID Number | Previous Passport Number | Previous ID Number |
| :-: | :--- | :---------------------: | :---------------: | :----------------------: | :----------------: |
| 1 | *AKILA NETHMINI* | N8074311 | 980260925V | \\- | \\- |
| 2 | *HISTORICAL DILANI* | P0262022 | 200477304008 | \\- | \\- |
| 3 | *SOON HIRUNI* | N11030886 | 820453352V | \\- | \\- |
| 4 | *NOBODY PIYUMI* | N0000001 | 999999999V | N0000002 | \\- |
| 5 | *MISMATCH KAVINDYA* | M2000001 | 933333333V | \\- | \\- |
| 6 | *NIC FALLBACK SEWWANDI* | SHORT | 922222222V | \\- | \\- |
| 7 | *SESSION FAIL THARUSHI* | E2000001 | 944444444V | \\- | \\- |
| 8 | *MALE LAKMAL BANDARA* | N7777777 | 877777777V | \\- | \\- |
| 9 | *MYSTERY PERSON* | N6666666 | 866666666V | \\- | \\- |
"""

GENDER_TXT = """| No. | Name | Current Passport Number | Current ID Number | Previous Passport Number | Previous ID Number | Gender |
| 1 | *FEMALE ONE* | N8074311 | 980260925V | \\- | \\- | Female |
| 2 | *MALE SKIP* | N1111111 | 811111111V | \\- | \\- | Male |
| 3 | *FEMALE TWO* | P0262022 | 200477304008 | \\- | \\- | F |
| 4 | *MALE TOKEN KAVINDU* | N2222222 | 822222222V | \\- | \\- | F |
"""


def test_date_logic():
    today = date(2026, 10, 23)
    from datetime import datetime
    expiry, days, state = fbr.expiry_info(datetime(2026, 1, 26, 15, 38, 52), today)
    assert expiry == date(2026, 7, 26) and state == "EXPIRED" and days < 0
    # month clamping: 31 Aug + 6 months -> 28 Feb
    assert fbr.add_months(date(2025, 8, 31), 6) == date(2026, 2, 28)
    # recent report is valid
    expiry, days, state = fbr.expiry_info(datetime.combine(today - timedelta(days=10),
                                                            datetime.min.time()), today)
    assert state == "VALID" and days > 14
    # 5 days left -> expires soon
    expiry, days, state = fbr.expiry_info(datetime.combine(today - timedelta(days=178),
                                                            datetime.min.time()), today)
    assert state == "EXPIRES SOON" and 0 <= days <= 14
    print("PASS test_date_logic")


def test_gender_filter():
    # explicit Gender column overrides the name heuristic (row 4 says F -> kept)
    p = WORK / "gender.txt"
    p.write_text(GENDER_TXT, encoding="utf-8")
    apps = fbr.load_applicants(str(p))
    assert [a.no for a in apps] == [1, 3, 4]

    # name-token filter on the main fixture: male row 8 and unknown row 9 gone
    (WORK / "name_female.txt").write_text(NAME_TXT, encoding="utf-8")
    apps2, skipped, uncertain = fbr.load_applicants_with_skips(str(WORK / "name_female.txt"))
    assert [a.no for a in apps2] == [1, 2, 3, 4, 5, 6, 7]
    assert any(n == 8 and "male" in why for n, _, why in skipped)
    assert [n for n, _ in uncertain] == [9]

    # direct classifier spot checks against the real 142-name vocabulary
    assert fbr.classify_gender("GALGODAPITIYA LAKMAL SANJEEWA BANDARA") == "M"
    assert fbr.classify_gender("HEMBATHANTHRIGE THARUSHIKA NADEESHANI") == "F"
    assert fbr.classify_gender("MYSTERY PERSON") == "?"
    print("PASS test_gender_filter")


def test_full_browser_flow():
    apps_file = WORK / "name_female.txt"
    apps_file.write_text(NAME_TXT, encoding="utf-8")

    # point the checker at the mock server and use an isolated session file
    fbr.FBR_URL = MOCK_FORM
    fbr.WELCOME_URL = MOCK_WELCOME
    fbr.SESSION_FILE = WORK / "session_fbr.json"
    fbr.NAV_TIMEOUT = 8000
    fbr.DELAY_BETWEEN = 0.05

    real_input = builtins.input
    builtins.input = lambda *a, **k: ""
    sys.argv = ["fbr_checker.py", "--names", str(apps_file),
                "--outdir", str(WORK / "output"), "--headless", "--delay", "0.02"]
    try:
        buf = io.StringIO()
        with redirect_stdout(buf):
            fbr.main()
        log = buf.getvalue()
    finally:
        builtins.input = real_input

    # the non-female rows were announced before the browser started
    assert "skipped 1 non-female" in log and "MALE LAKMAL" in log
    assert "cannot be gender-identified" in log and "MYSTERY PERSON" in log
    # the menu walk reached the FBR page through the welcome-page link
    assert "via portal menu" in log

    summary = json.loads((WORK / "output" / "fbr_summary.json").read_text(encoding="utf-8"))
    by_no = {r["no"]: r for r in summary["results"]}
    counts = summary["counts"]
    assert summary["females_checked"] == 7
    assert summary["applicants_in_file"] == 9
    assert len(summary["skipped_non_female"]) == 1
    assert summary["uncertain_names"][0]["no"] == 9

    # 1 VALID (recent Recommended report, found by passport)
    r1 = by_no[1]
    assert r1["expiry_state"] == "VALID" and r1["validation"] == "OK"
    assert r1["fbr_status"] == "Recommended" and r1["match_on"] == "Passport+NIC"
    assert r1["result_passport"] == "N8074311" and r1["result_nic"] == "980260925V"

    # 2 EXPIRED (fixed 26/01/2026 report)
    r2 = by_no[2]
    assert r2["expiry_state"] == "EXPIRED" and r2["expiry_date"] == "26/07/2026"
    assert "expired" in r2["remark"].lower()

    # 3 EXPIRES SOON
    assert by_no[3]["expiry_state"] == "EXPIRES SOON"

    # 4 NOT FOUND under passport + NIC
    r4 = by_no[4]
    assert r4["validation"] == "NOT FOUND" and r4["expiry_state"] == ""

    # 5 MISMATCH (row belongs to different PP/NIC)
    r5 = by_no[5]
    assert r5["validation"] == "MISMATCH"
    assert r5["result_passport"] == "N9999999" and r5["result_nic"] == "111111111V"

    # 6 short passport alert handled, NIC fallback finds the row
    r6 = by_no[6]
    assert r6["searched_by"] == "NIC" and r6["result_nic"] == "922222222V"
    assert r6["validation"] in ("OK", "CHECK")

    # 7 session expiry on every attempt -> ERROR row, run does not crash
    r7 = by_no[7]
    assert r7["validation"] == "ERROR"
    assert "re-login" in r7["remark"].lower()

    # evidence screenshots exist for every row actually found
    shots = list((WORK / "output" / "fbr_screenshots").glob("*.png"))
    assert len(shots) >= 4 and all(s.stat().st_size > 1000 for s in shots)

    # CSV + Excel + latest copies
    assert (WORK / "output" / "fbr_report_latest.csv").exists()
    assert (WORK / "output" / "fbr_report_latest.xlsx").exists()
    with open(WORK / "output" / "fbr_report_latest.csv", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 7 and "Expiry State" in rows[0]

    # one workbook with one sheet per outcome group
    wb = load_workbook(WORK / "output" / "fbr_report_latest.xlsx")
    assert wb.sheetnames == ["All Females", "VALID FBR", "EXPIRED & EXPIRING",
                              "NO FBR", "REVIEW & ERRORS"]
    assert wb["All Females"].max_row == 8           # header + 7 females
    assert wb["VALID FBR"].max_row == 3             # #1 + #6 (NIC fallback)
    assert wb["EXPIRED & EXPIRING"].max_row == 3    # #2 + #3
    assert wb["NO FBR"].max_row == 2                # only #4
    assert wb["REVIEW & ERRORS"].max_row == 3       # #5 MISMATCH + #7 ERROR
    # the groups are mutually exclusive and together cover every female
    grouped = sum(wb[s].max_row - 1 for s in ("VALID FBR", "EXPIRED & EXPIRING",
                                              "NO FBR", "REVIEW & ERRORS"))
    assert grouped == 7

    # console summary prints the main buckets and named lists
    for token in ("VALID FBR", "EXPIRED", "NOT FOUND (no FBR)", "GOT VALID FBR"):
        assert token in log
    print("PASS test_full_browser_flow", counts)


def main():
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)

    server = ThreadingHTTPServer(("127.0.0.1", 8766), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        test_date_logic()
        test_gender_filter()
        test_full_browser_flow()
    finally:
        server.shutdown()
    print("\nAll System 2 (FBR) tests passed.")


if __name__ == "__main__":
    main()
