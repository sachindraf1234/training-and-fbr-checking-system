"""
End-to-end test for slbfe_training_checker.py against the local mock SLBFE
server. Covers:
  * TRAINED detection via the "Romania Pre departure Orientation Programme"
    marker + proof screenshot
  * NOT_TRAINED / NOT_FOUND (yellow Not Found page + per-number screenshot) /
    REVIEW / INCOMPLETE / INVALID_FORMAT / blacklist branches
  * a single training_results.xlsx workbook (Applicants + Search Log sheets)
  * progress.json skip of TRAINED people, --recheck-all, --only ranges,
    e-mail no-op, and session-expiry handling.
No real SLBFE credentials are used.
"""
import builtins
import io
import json
import shutil
import sys
import threading
from contextlib import redirect_stdout
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import slbfe_training_checker as m
from mock_slbfe_server import Handler, EXPIRED_ID

MOCK_URL = "http://127.0.0.1:8765/TraRegChk.asp"
WORK = ROOT / "tests" / "v2_workspace"

# Markdown table in the same shape as the real name.txt
NAME_TXT = """| No. | Name | Current Passport Number | Current ID Number | Previous Passport Number | Previous ID Number |
| :-: | :--- | :---------------------: | :---------------: | :----------------------: | :----------------: |
| 1 | *TRAINED PERSON* | N8074311 | 980260925V | \\- | \\- |
| 2 | *ORIENTATION ONLY* | N5408869 | 820453352V | \\- | \\- |
| 3 | *REGISTRATION ONLY* | R1234567 | 900000000V | \\- | \\- |
| 4 | *NO RECORDS AT ALL* | N0000001 | 999999999V | N0000002 | \\- |
| 5 | *REVIEW MISMATCH* | M2000001 | 922222222V | \\- | \\- |
| 6 | *BLACKLISTED WORKER* | B1000001 | 911111111V | \\- | \\- |
| 7 | *BAD FORMAT THEN NIC* | SHORT | 933333333V | \\- | \\- |
| 8 | *SESSION EXPIRED CASE* | E2000001 | 944444444V | \\- | \\- |
"""

EXPECTED = {
    1: "TRAINED",
    2: "NOT_TRAINED",
    3: "NOT_TRAINED",
    4: "NOT_FOUND",
    5: "REVIEW",
    6: "NOT_TRAINED",
    7: "NOT_FOUND",
    8: "INCOMPLETE",
}


def run_cli(*extra):
    sys.argv = ["slbfe_training_checker.py",
                "--input", str(WORK / "name.txt"),
                "--outdir", str(WORK),
                "--headless", "--delay", "0.05", "--timeout", "8", *extra]
    buf = io.StringIO()
    with redirect_stdout(buf):
        m.main()
    return buf.getvalue()


def sheet_rows(sheet):
    hdr = [c.value for c in sheet[1]]
    return [dict(zip(hdr, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]


def read_workbook():
    wb = load_workbook(WORK / "training_results.xlsx")
    assert wb.sheetnames == ["Applicants", "Search Log"]
    applicants = {int(r["no"]): r for r in sheet_rows(wb["Applicants"])}
    log = sheet_rows(wb["Search Log"])
    return applicants, log


def main():
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)
    (WORK / "name.txt").write_text(NAME_TXT, encoding="utf-8")

    server = ThreadingHTTPServer(("127.0.0.1", 8765), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    m.INQUIRY_URL = MOCK_URL
    m.PORTAL_URL = MOCK_URL
    m.SESSION_FILE = WORK / "session.json"   # keep the real system folder clean

    real_input = builtins.input
    builtins.input = lambda *a, **k: ""   # auto-press ENTER at login prompts
    try:
        # ---------------------------------------------------------- full run
        out = run_cli()
        assert "8 applicants loaded" in out
        summary, log = read_workbook()
        for no, status in EXPECTED.items():
            assert summary[no]["status"] == status, (no, summary[no]["status"])

        # #1 TRAINED details + proof screenshot
        t1 = summary[1]
        assert t1["found_by"] == "Current Passport=N8074311"
        assert t1["training_date"] == "14/10/2026"
        assert "Romania Pre departure Orientation Programme" in t1["training"]
        trained_shot = WORK / "screenshots" / "001_TRAINED_PERSON_TRAINED.png"
        assert trained_shot.exists() and trained_shot.stat().st_size > 1000
        assert trained_shot.name in t1["screenshots"]

        # #4 NOT_FOUND: all three numbers got yellow Not Found pages -> 3 screenshots
        nf4 = sorted(p.name for p in (WORK / "screenshots").glob("004_*.png"))
        assert nf4 == [
            "004_NO_RECORDS_AT_ALL_CurrentNIC_999999999V.png",
            "004_NO_RECORDS_AT_ALL_CurrentPassport_N0000001.png",
            "004_NO_RECORDS_AT_ALL_PreviousPassport_N0000002.png",
        ]
        assert "Not Found page(s)" in summary[4]["note"]

        # #7 invalid short passport recorded, NIC came back Not Found + screenshot
        l7 = [r for r in log if int(r["no"]) == 7]
        assert any(r["status"] == "INVALID_FORMAT" and r["id_searched"] == "SHORT" for r in l7)
        assert any(r["status"] == "NOT_FOUND" and r["id_searched"] == "933333333V" for r in l7)
        assert (WORK / "screenshots" / "007_BAD_FORMAT_THEN_NIC_CurrentNIC_933333333V.png").exists()

        # #8 session-expiry recorded in the Search Log -> applicant INCOMPLETE
        assert any(int(r["no"]) == 8 and r["status"] == "SESSION_EXPIRED" for r in log)

        # #5 REVIEW note explains the PP/NIC mismatch
        assert "not this applicant" in summary[5]["note"]

        # #6 blacklist flag captured even though not orientation-trained
        assert "Blacklisted" in summary[6]["blacklisted"]

        # progress.json persisted after every applicant
        progress = json.loads((WORK / "progress.json").read_text(encoding="utf-8"))
        assert len(progress) == 8
        assert progress["N8074311"]["status"] == "TRAINED"

        # exactly ONE Excel workbook; old csv/txt/zip artifacts must not exist
        assert (WORK / "training_results.xlsx").exists()
        for stale in ("results_summary.csv", "not_trained.csv", "not_trained.txt",
                      "results_detail.csv", "screenshots.zip"):
            assert not (WORK / stale).exists(), stale
        # no unknown pages -> no debug dumps
        assert not (WORK / "debug").exists()

        # every yellow Not Found page is screenshotted (one per unfiled number),
        # plus the single TRAINED proof -> 10 files total
        shots = sorted(p.name for p in (WORK / "screenshots").glob("*.png"))
        assert shots == [
            "001_TRAINED_PERSON_TRAINED.png",
            "002_ORIENTATION_ONLY_CurrentNIC_820453352V.png",
            "003_REGISTRATION_ONLY_CurrentNIC_900000000V.png",
            "004_NO_RECORDS_AT_ALL_CurrentNIC_999999999V.png",
            "004_NO_RECORDS_AT_ALL_CurrentPassport_N0000001.png",
            "004_NO_RECORDS_AT_ALL_PreviousPassport_N0000002.png",
            "005_REVIEW_MISMATCH_CurrentNIC_922222222V.png",
            "006_BLACKLISTED_WORKER_CurrentNIC_911111111V.png",
            "007_BAD_FORMAT_THEN_NIC_CurrentNIC_933333333V.png",
            "008_SESSION_EXPIRED_CASE_CurrentNIC_944444444V.png",
        ], shots

        # --------------------------------------------- second run skips TRAINED
        out2 = run_cli()
        assert "1 already TRAINED (skipped) | 7 to check" in out2
        # same filenames are overwritten, count stays stable
        assert len(list((WORK / "screenshots").glob("*.png"))) == 10

        # --------------------------------------------- --only with a range
        out3 = run_cli("--only", "1-2", "--recheck-all")
        assert "2 applicants loaded" in out3

        # --------------------------------------------- e-mail no-op path
        out5 = run_cli("--only", "2", "--recheck-all", "--notify-email")
        assert "e-mail not sent" in out5
    finally:
        builtins.input = real_input
        server.shutdown()

    print("\nAll v2 end-to-end assertions passed.")


if __name__ == "__main__":
    main()
